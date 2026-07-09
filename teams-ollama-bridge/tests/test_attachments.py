"""Tests für Attachment-Verarbeitung."""

from __future__ import annotations

import json
from pathlib import Path

import httpx
import pytest
import respx
from tests.conftest import write_input_file

from teams_ollama_bridge.attachment_resolver import AttachmentResolver
from teams_ollama_bridge.attachment_service import AttachmentService
from teams_ollama_bridge.document_extractor import DocumentExtractor
from teams_ollama_bridge.exceptions import AttachmentNotSyncedError, AttachmentPathError
from teams_ollama_bridge.file_service import is_file_stable
from teams_ollama_bridge.logging_config import setup_logging
from teams_ollama_bridge.mock_processor import MockProcessor
from teams_ollama_bridge.models import AttachmentInfo, ImageProcessingMode, InputRequest
from teams_ollama_bridge.ollama_client import OllamaClient
from teams_ollama_bridge.processor import RequestProcessor
from teams_ollama_bridge.repository import RequestRepository


def _resolver(settings, input_dir: Path) -> AttachmentResolver:
    return AttachmentResolver(
        input_dir=input_dir,
        attachments_base_dir=input_dir,
        allowed_extensions=settings.parsed_allowed_extensions,
        max_file_size_bytes=settings.attachments_max_file_size_bytes,
        max_files=settings.attachments_max_files,
    )


def test_request_without_attachments_still_works(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    write_input_file(settings.input_dir, request_id="text-only")
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    assert processor.process_file(list(settings.input_dir.glob("request_*.json"))[0]) is True


def test_empty_attachments_array(settings) -> None:
    request = InputRequest(
        requestId="a1",
        messageId="m",
        chatId="c",
        message="Hallo",
        attachments=[],
    )
    assert request.attachments == []


def test_txt_attachment_extracted(settings, workspace) -> None:
    files_dir = settings.input_dir / "files"
    txt_path = files_dir / "att-001_hallo.txt"
    txt_path.write_text("Hallo aus TXT", encoding="utf-8")
    extractor = DocumentExtractor()
    content = extractor.extract(txt_path, ".txt", 1000)
    assert "Hallo aus TXT" in content


def test_utf8_bom_text_file(settings, workspace) -> None:
    files_dir = settings.input_dir / "files"
    path = files_dir / "bom.txt"
    path.write_bytes(b"\xef\xbb\xbfUTF-8 mit BOM")
    extractor = DocumentExtractor()
    assert "UTF-8 mit BOM" in extractor.extract(path, ".txt", 1000)


def test_pdf_attachment_extracted(settings, workspace) -> None:
    import fitz

    files_dir = settings.input_dir / "files"
    pdf_path = files_dir / "doc.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "PDF Testinhalt")
    doc.save(pdf_path)
    doc.close()
    extractor = DocumentExtractor()
    content = extractor.extract(pdf_path, ".pdf", 5000)
    assert "PDF Testinhalt" in content
    assert "[Seite 1]" in content


def test_docx_attachment_extracted(settings, workspace) -> None:
    from docx import Document

    files_dir = settings.input_dir / "files"
    path = files_dir / "doc.docx"
    doc = Document()
    doc.add_paragraph("DOCX Absatz")
    doc.save(path)
    extractor = DocumentExtractor()
    assert "DOCX Absatz" in extractor.extract(path, ".docx", 5000)


def test_xlsx_attachment_extracted(settings, workspace) -> None:
    from openpyxl import Workbook

    files_dir = settings.input_dir / "files"
    path = files_dir / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Daten"
    ws["A1"] = "Wert"
    wb.save(path)
    extractor = DocumentExtractor()
    content = extractor.extract(path, ".xlsx", 5000)
    assert "Wert" in content


def test_not_copied_attachment_does_not_abort(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    write_input_file(
        settings.input_dir,
        request_id="not-copied",
        message="Bitte zusammenfassen",
        extra={
            "attachments": [
                {
                    "name": "fehlt.pdf",
                    "localPath": "",
                    "status": "not_copied",
                    "error": "Flow konnte nicht kopieren",
                }
            ]
        },
    )
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    assert processor.process_file(list(settings.input_dir.glob("request_*.json"))[0]) is True
    output = json.loads(
        (settings.output_dir / "response_not-copied.json").read_text(encoding="utf-8")
    )
    assert output["status"] == "completed"
    assert "attachmentsProcessed" in output


def test_missing_attachment_temporary_error(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    write_input_file(
        settings.input_dir,
        request_id="missing-file",
        message="Test",
        extra={
            "attachments": [
                {"name": "fehlt.pdf", "localPath": "files/fehlt.pdf"}
            ]
        },
    )
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    path = list(settings.input_dir.glob("request_*.json"))[0]
    assert processor.process_file(path) is False
    assert path.exists()


def test_path_traversal_blocked(settings, workspace) -> None:
    resolver = _resolver(settings, settings.input_dir)
    with pytest.raises(AttachmentPathError):
        resolver.resolve_local_path("../secret.txt")


def test_absolute_path_blocked(settings, workspace) -> None:
    resolver = _resolver(settings, settings.input_dir)
    with pytest.raises(AttachmentPathError):
        resolver.resolve_local_path("C:\\Windows\\system.ini")


def test_disallowed_extension_skipped(settings, workspace) -> None:
    files_dir = settings.input_dir / "files"
    (files_dir / "bad.exe").write_bytes(b"MZ")
    service = AttachmentService(settings)
    request = InputRequest(
        requestId="ext",
        messageId="m",
        chatId="c",
        message="Test",
        attachments=[AttachmentInfo(name="bad.exe", localPath="files/bad.exe")],
    )
    batch = service.process_request(request, treat_missing_as_failed=True)
    assert batch.processed[0].status.value == "failed"


def test_image_metadata_mode(settings, workspace) -> None:
    from PIL import Image

    files_dir = settings.input_dir / "files"
    img_path = files_dir / "bild.png"
    Image.new("RGB", (100, 50), color="red").save(img_path)
    service = AttachmentService(settings)
    request = InputRequest(
        requestId="img",
        messageId="m",
        chatId="c",
        message="Was ist auf dem Bild?",
        attachments=[AttachmentInfo(name="bild.png", localPath="files/bild.png")],
    )
    batch = service.process_request(request)
    assert batch.processed[0].status.value == "processed"
    assert "100x50" in batch.processed[0].prompt_section


@respx.mock
def test_ollama_vision_mocked(settings, workspace) -> None:
    from PIL import Image

    settings.image_processing_mode = ImageProcessingMode.OLLAMA_VISION
    files_dir = settings.input_dir / "files"
    img_path = files_dir / "vision.png"
    Image.new("RGB", (20, 20), color="blue").save(img_path)

    client = OllamaClient(
        base_url="http://127.0.0.1:11434",
        model="test",
        timeout_seconds=5,
        keep_alive="10m",
        temperature=0.2,
        system_prompt="Test",
        max_output_characters=1000,
    )
    respx.post("http://127.0.0.1:11434/api/chat").mock(
        return_value=httpx.Response(
            200,
            json={"message": {"content": "Ein blaues Quadrat"}},
        )
    )
    service = AttachmentService(settings, client)
    request = InputRequest(
        requestId="vis",
        messageId="m",
        chatId="c",
        message="Bild?",
        attachments=[AttachmentInfo(name="vision.png", localPath="files/vision.png")],
    )
    batch = service.process_request(request)
    assert "blaues" in batch.processed[0].prompt_section.lower()


def test_mock_processor_with_attachments(settings, workspace) -> None:
    files_dir = settings.input_dir / "files"
    (files_dir / "a.txt").write_text("Inhalt", encoding="utf-8")
    service = AttachmentService(settings)
    request = InputRequest(
        requestId="mock-att",
        messageId="m",
        chatId="c",
        message="Zusammenfassung",
        attachments=[AttachmentInfo(name="a.txt", localPath="files/a.txt")],
    )
    batch = service.process_request(request)
    mock = MockProcessor(20000)
    result = mock.process("Zusammenfassung", batch)
    assert "Erkannte Anhänge" in result.answer
    assert "a.txt" in result.answer


def test_output_contains_attachments_processed(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    files_dir = settings.input_dir / "files"
    (files_dir / "att-002_data.txt").write_text("Dateninhalt", encoding="utf-8")
    write_input_file(
        settings.input_dir,
        request_id="att-002",
        message="Fasse zusammen",
        extra={
            "attachments": [
                {
                    "name": "data.txt",
                    "localPath": "files/att-002_data.txt",
                }
            ]
        },
    )
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    processor.process_file(list(settings.input_dir.glob("request_*.json"))[0])
    data = json.loads((settings.output_dir / "response_att-002.json").read_text(encoding="utf-8"))
    assert data["status"] == "completed"
    assert "answer" in data
    assert data["attachmentsProcessed"][0]["status"] == "processed"


def test_attachments_archived_after_success(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    files_dir = settings.input_dir / "files"
    att = files_dir / "att-003_file.txt"
    att.write_text("Archivtest", encoding="utf-8")
    write_input_file(
        settings.input_dir,
        request_id="att-003",
        message="Test",
        extra={"attachments": [{"name": "file.txt", "localPath": "files/att-003_file.txt"}]},
    )
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    processor.process_file(list(settings.input_dir.glob("request_*.json"))[0])
    assert not att.exists()
    archived = list((settings.processed_input_dir / "files").glob("*.txt"))
    assert archived


def test_unstable_attachment_not_processed(settings, workspace) -> None:
    files_dir = settings.input_dir / "files"
    att = files_dir / "unstable.txt"
    att.write_text("x", encoding="utf-8")
    service = AttachmentService(settings)
    settings.file_stable_seconds = 60.0
    request = InputRequest(
        requestId="u",
        messageId="m",
        chatId="c",
        message="Test",
        attachments=[AttachmentInfo(name="u.txt", localPath="files/unstable.txt")],
    )
    with pytest.raises(AttachmentNotSyncedError):
        service.check_attachments_stable(request)
    assert is_file_stable(att, 60.0) is False


def test_repository_attachment_migration(settings, workspace) -> None:
    repo = RequestRepository(settings.database_path)
    repo.save_attachments("req-1", [])
    assert repo.list_attachments("req-1") == []


def test_file_too_large_skipped(settings, workspace) -> None:
    files_dir = settings.input_dir / "files"
    (files_dir / "big.txt").write_text("x" * 100, encoding="utf-8")
    settings.attachments_max_file_size_mb = 0
    service = AttachmentService(settings)
    request = InputRequest(
        requestId="big",
        messageId="m",
        chatId="c",
        message="Test",
        attachments=[AttachmentInfo(name="big.txt", localPath="files/big.txt")],
    )
    batch = service.process_request(request, treat_missing_as_failed=True)
    assert batch.processed[0].status.value == "failed"
    assert batch.processed[0].error is not None
    assert "Größe" in batch.processed[0].error


def test_empty_pdf_hint(settings, workspace) -> None:
    import fitz

    files_dir = settings.input_dir / "files"
    pdf_path = files_dir / "empty.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()
    service = AttachmentService(settings)
    request = InputRequest(
        requestId="empty-pdf",
        messageId="m",
        chatId="c",
        message="Test",
        attachments=[AttachmentInfo(name="empty.pdf", localPath="files/empty.pdf")],
    )
    batch = service.process_request(request)
    assert batch.processed[0].status.value == "failed"
    assert "Kein extrahierbarer Text" in (batch.processed[0].error or "")


def test_encrypted_pdf_error(settings, workspace) -> None:
    import fitz

    files_dir = settings.input_dir / "files"
    pdf_path = files_dir / "enc.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Secret")
    doc.save(pdf_path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret")
    doc.close()
    service = AttachmentService(settings)
    request = InputRequest(
        requestId="enc-pdf",
        messageId="m",
        chatId="c",
        message="Test",
        attachments=[AttachmentInfo(name="enc.pdf", localPath="files/enc.pdf")],
    )
    batch = service.process_request(request)
    assert batch.processed[0].status.value == "failed"
    assert batch.processed[0].error is not None
    assert "verschlüsselt" in batch.processed[0].error.lower()


def test_path_outside_input_dir_blocked(settings, workspace) -> None:
    outside = workspace / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    resolver = AttachmentResolver(
        input_dir=settings.input_dir,
        attachments_base_dir=settings.input_dir,
        allowed_extensions=settings.parsed_allowed_extensions,
        max_file_size_bytes=settings.attachments_max_file_size_bytes,
        max_files=settings.attachments_max_files,
    )
    with pytest.raises(AttachmentPathError):
        resolver.resolve_local_path("../outside.txt")


def test_output_flow2_compatible(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    write_input_file(settings.input_dir, request_id="flow2", message="Hallo")
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    processor.process_file(list(settings.input_dir.glob("request_*.json"))[0])
    data = json.loads((settings.output_dir / "response_flow2.json").read_text(encoding="utf-8"))
    for key in ("requestId", "messageId", "chatId", "answer", "status", "processedAt"):
        assert key in data
    assert data["status"] == "completed"
    assert isinstance(data["answer"], str)


def test_missing_file_after_retry_limit(settings, workspace) -> None:
    setup_logging("WARNING", settings.log_file_path, 10000, 1)
    write_input_file(
        settings.input_dir,
        request_id="miss-limit",
        message="Test",
        extra={"attachments": [{"name": "gone.txt", "localPath": "files/gone.txt"}]},
    )
    settings.max_process_retries = 1
    repo = RequestRepository(settings.database_path)
    processor = RequestProcessor(settings, repo)
    path = list(settings.input_dir.glob("request_*.json"))[0]
    assert processor.process_file(path) is False
    assert path.exists()
    assert processor.process_file(path) is True
    output_path = settings.output_dir / "response_miss-limit.json"
    data = json.loads(output_path.read_text(encoding="utf-8"))
    assert data["status"] == "completed"
    assert data["attachmentsProcessed"][0]["status"] == "failed"
